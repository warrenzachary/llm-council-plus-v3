"""
Plain-text extraction from uploaded documents.

Each extractor is wrapped in try/except and returns a safe fallback
on failure. Output is truncated to max_chars so a single file
cannot blow up the LLM context window.
"""

from pathlib import Path

# Extensions that should be read as plain text (UTF-8, errors ignored)
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".css", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".log", ".sh", ".bat", ".ps1",
}


def extract_text(path: Path, max_chars: int = 50_000) -> str:
    """
    Return plain text from `path`, truncated to `max_chars`.

    Never raises. On failure returns a bracketed error message.
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".pptx":
            text = _extract_pptx(path)
        elif suffix == ".xlsx":
            text = _extract_xlsx(path)
        elif suffix in _TEXT_EXTENSIONS:
            text = _extract_plaintext(path)
        else:
            # Unknown binary format, do not read garbled bytes
            return f"[Unsupported file type: {suffix}]"

        return (text or "")[:max_chars]
    except Exception as exc:
        return f"[Could not extract text from {path.name}: {exc}]"


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            parts.append(page_text)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = (paragraph.text or "").strip()
                    if text:
                        slide_texts.append(text)
        if slide_texts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            rows.append("\t".join(cells))
        if rows:
            parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _extract_plaintext(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore") as f:
        return f.read()
